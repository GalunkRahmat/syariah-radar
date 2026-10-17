"""
Historical Screening Database & Excel Export Module
Stores daily screening snapshots in SQLite for trend/behavior analysis and exports to Excel (.xlsx).
"""

import io
import os
import shutil
import json
import sqlite3
from datetime import datetime
from pathlib import Path
from typing import Dict, Any, List, Optional
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

PROJECT_ROOT = Path(__file__).resolve().parent.parent
if os.environ.get("VERCEL"):
    DATA_DIR = Path("/tmp/data")
else:
    DATA_DIR = PROJECT_ROOT / "data"
DB_FILE = DATA_DIR / "screening_history.db"


def get_db_connection() -> sqlite3.Connection:
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    seed_db = PROJECT_ROOT / "data" / "screening_history.db"
    if os.environ.get("VERCEL") and not DB_FILE.exists() and seed_db.exists():
        try:
            shutil.copy2(seed_db, DB_FILE)
        except Exception:
            pass
    conn = sqlite3.connect(DB_FILE)
    conn.row_factory = sqlite3.Row
    return conn


def init_history_db() -> None:
    """Initializes SQLite schema for historical market tracking."""
    conn = get_db_connection()
    cursor = conn.cursor()

    # Table for daily screening run metadata
    cursor.execute("""
        CREATE TABLE IF NOT EXISTS daily_snapshots (
            market_date TEXT PRIMARY KEY,
            total_scanned INTEGER,
            total_candidates INTEGER,
            summary_json TEXT,
            created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
        )
    """)

    # Table for individual stock candidates per day
    cursor.execute("""
        CREATE TABLE IF NOT EXISTS candidate_history (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            market_date TEXT,
            code TEXT,
            name TEXT,
            sector TEXT,
            score INTEGER,
            setup_type TEXT,
            risk_level TEXT,
            close_price REAL,
            entry_price REAL,
            stop_loss REAL,
            tp1 REAL,
            tp2 REAL,
            rr_ratio TEXT,
            bandar_status TEXT,
            vol_ratio REAL,
            rsi REAL,
            mfi REAL,
            der REAL,
            roe REAL,
            details_json TEXT,
            created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
            UNIQUE(market_date, code)
        )
    """)

    cursor.execute("CREATE INDEX IF NOT EXISTS idx_cand_date ON candidate_history(market_date)")
    cursor.execute("CREATE INDEX IF NOT EXISTS idx_cand_code ON candidate_history(code)")

    conn.commit()
    conn.close()


def save_screening_snapshot(market_date: str, candidates: List[Dict[str, Any]], summary: Dict[str, Any]) -> None:
    """Saves or updates daily screening snapshot and all candidate details."""
    init_history_db()
    conn = get_db_connection()
    cursor = conn.cursor()

    try:
        # Upsert daily snapshot summary
        cursor.execute("""
            INSERT INTO daily_snapshots (market_date, total_scanned, total_candidates, summary_json, created_at)
            VALUES (?, ?, ?, ?, CURRENT_TIMESTAMP)
            ON CONFLICT(market_date) DO UPDATE SET
                total_scanned=excluded.total_scanned,
                total_candidates=excluded.total_candidates,
                summary_json=excluded.summary_json,
                created_at=CURRENT_TIMESTAMP
        """, (
            market_date,
            summary.get("total_scanned", 0),
            summary.get("total_candidates", len(candidates)),
            json.dumps(summary)
        ))

        # Insert or replace candidate records
        for c in candidates:
            code = c.get("code")
            tp = c.get("trading_plan", {})
            bandar = c.get("bandarmologi", {})
            fund = c.get("fundamental") or c.get("fundamentals") or {}

            close_price = float(c.get("price") or c.get("close") or tp.get("entry", 0.0))
            rsi_val = float(c.get("rsi", 50.0))
            mfi_val = float(c.get("mfi", 50.0))
            vol_ratio_val = float(c.get("vol_ratio") or bandar.get("vol_ratio", 1.0))
            der_val = float(fund.get("der", 0.0) or 0.0)
            roe_val = float(fund.get("roe", 0.0) or 0.0)

            cursor.execute("""
                INSERT INTO candidate_history (
                    market_date, code, name, sector, score, setup_type, risk_level,
                    close_price, entry_price, stop_loss, tp1, tp2, rr_ratio,
                    bandar_status, vol_ratio, rsi, mfi, der, roe, details_json, created_at
                )
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, CURRENT_TIMESTAMP)
                ON CONFLICT(market_date, code) DO UPDATE SET
                    name=excluded.name,
                    sector=excluded.sector,
                    score=excluded.score,
                    setup_type=excluded.setup_type,
                    risk_level=excluded.risk_level,
                    close_price=excluded.close_price,
                    entry_price=excluded.entry_price,
                    stop_loss=excluded.stop_loss,
                    tp1=excluded.tp1,
                    tp2=excluded.tp2,
                    rr_ratio=excluded.rr_ratio,
                    bandar_status=excluded.bandar_status,
                    vol_ratio=excluded.vol_ratio,
                    rsi=excluded.rsi,
                    mfi=excluded.mfi,
                    der=excluded.der,
                    roe=excluded.roe,
                    details_json=excluded.details_json,
                    created_at=CURRENT_TIMESTAMP
            """, (
                market_date,
                code,
                c.get("name", ""),
                c.get("sector", ""),
                c.get("score", 0),
                c.get("setup_name") or c.get("setup_label") or c.get("setup_type", ""),
                c.get("risk_label") or c.get("risk_level", ""),
                close_price,
                float(tp.get("entry", close_price)),
                float(tp.get("stop_loss", 0.0)),
                float(tp.get("tp1", 0.0)),
                float(tp.get("tp2", 0.0)),
                tp.get("risk_reward") or tp.get("rr_ratio", "1 : 2.5"),
                bandar.get("label") or bandar.get("status_label") or bandar.get("status", "NEUTRAL"),
                vol_ratio_val,
                rsi_val,
                mfi_val,
                der_val,
                roe_val,
                json.dumps(c)
            ))

        conn.commit()
    finally:
        conn.close()


def get_history_dates() -> List[Dict[str, Any]]:
    """Returns list of available screening dates."""
    init_history_db()
    conn = get_db_connection()
    cursor = conn.cursor()
    cursor.execute("""
        SELECT market_date, total_scanned, total_candidates, summary_json, created_at
        FROM daily_snapshots
        ORDER BY market_date DESC
    """)
    rows = cursor.fetchall()
    dates = []
    for r in rows:
        dates.append({
            "market_date": r["market_date"],
            "total_scanned": r["total_scanned"],
            "total_candidates": r["total_candidates"],
            "created_at": r["created_at"],
            "summary": json.loads(r["summary_json"]) if r["summary_json"] else {}
        })
    conn.close()
    return dates


def get_history_candidates(market_date: str) -> Optional[Dict[str, Any]]:
    """Retrieves full candidate data for a specific historical date."""
    init_history_db()
    conn = get_db_connection()
    cursor = conn.cursor()

    cursor.execute("SELECT * FROM daily_snapshots WHERE market_date = ?", (market_date,))
    snap = cursor.fetchone()
    if not snap:
        conn.close()
        return None

    cursor.execute("SELECT details_json FROM candidate_history WHERE market_date = ? ORDER BY score DESC", (market_date,))
    rows = cursor.fetchall()
    candidates = [json.loads(r["details_json"]) for r in rows]

    summary = json.loads(snap["summary_json"]) if snap["summary_json"] else {
        "total_scanned": snap["total_scanned"],
        "total_candidates": snap["total_candidates"]
    }

    conn.close()
    return {"market_date": market_date, "summary": summary, "candidates": candidates}


def get_stock_timeline(code: str) -> List[Dict[str, Any]]:
    """Retrieves signal timeline for a single stock across history."""
    init_history_db()
    conn = get_db_connection()
    cursor = conn.cursor()
    cursor.execute("""
        SELECT market_date, score, setup_type, risk_level, close_price, entry_price, stop_loss, tp1, tp2, bandar_status, vol_ratio, rsi
        FROM candidate_history
        WHERE code = ?
        ORDER BY market_date ASC
    """, (code.upper(),))
    rows = cursor.fetchall()
    timeline = [dict(r) for r in rows]
    conn.close()
    return timeline


def generate_excel_export(market_date: str, candidates: List[Dict[str, Any]], summary: Optional[Dict[str, Any]] = None) -> io.BytesIO:
    """Generates professionally styled Excel (.xlsx) workbook from candidates list."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"Radar Saham {market_date}"

    # Modern Palette Styling
    header_fill = PatternFill(start_color="10B981", end_color="10B981", fill_type="solid") # Emerald green
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    title_font = Font(name="Calibri", size=14, bold=True, color="0F172A")
    subtitle_font = Font(name="Calibri", size=10, italic=True, color="64748B")
    bold_font = Font(name="Calibri", size=10, bold=True)
    normal_font = Font(name="Calibri", size=10)
    
    thin_border = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0')
    )

    # Title Block
    ws['A1'] = "LAPORAN HASIL SCREENING SWING SAHAM SYARIAH BEI (ISSI / JII / JII70)"
    ws['A1'].font = title_font
    ws['A2'] = f"Tanggal Data Pasar: {market_date} | Total Kandidat Lolos: {len(candidates)} Saham | Sistem Validasi 3-Pilar"
    ws['A2'].font = subtitle_font

    headers = [
        "No",
        "Kode Saham",
        "Nama Perusahaan",
        "Sektor",
        "Skor Total",
        "Kategori Setup",
        "Tingkat Risiko",
        "Harga Terakhir (Rp)",
        "Area Beli Entry (Rp)",
        "Stop Loss (Rp)",
        "Potensi Risiko SL (%)",
        "Target Profit 1 (Rp)",
        "Estimasi Cuan TP1 (%)",
        "Target Profit 2 (Rp)",
        "Estimasi Cuan TP2 (%)",
        "Risk:Reward",
        "Status Bandar",
        "Volume Ratio",
        "RSI 14",
        "MFI 14",
        "Stoch RSI",
        "MACD Status",
        "DER (%)",
        "ROE (%)"
    ]

    header_row = 4
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=header_row, column=col_idx, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border
    ws.row_dimensions[header_row].height = 28

    # Zebra striping fills
    zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

    for i, c in enumerate(candidates, start=1):
        row_idx = header_row + i
        tp = c.get("trading_plan", {})
        bandar = c.get("bandarmologi", {})
        fund = c.get("fundamental") or c.get("fundamentals") or {}

        close_price = int(round(float(c.get("price") or c.get("close") or tp.get("entry", 0))))
        entry = int(round(float(tp.get("entry", close_price))))
        sl = int(round(float(tp.get("stop_loss", 0))))
        tp1 = int(round(float(tp.get("tp1", 0))))
        tp2 = int(round(float(tp.get("tp2", 0))))

        sl_pct = tp.get("risk_pct")
        if sl_pct is None or sl_pct == 0:
            sl_pct = round(abs((entry - sl) / entry) * 100, 2) if entry > 0 else 0
        else:
            sl_pct = round(abs(float(sl_pct)), 2)

        tp1_pct = tp.get("tp1_pct")
        if tp1_pct is None or tp1_pct == 0:
            tp1_pct = round(((tp1 - entry) / entry) * 100, 2) if entry > 0 else 0
        else:
            tp1_pct = round(float(tp1_pct), 2)

        tp2_pct = tp.get("tp2_pct")
        if tp2_pct is None or tp2_pct == 0:
            tp2_pct = round(((tp2 - entry) / entry) * 100, 2) if entry > 0 else 0
        else:
            tp2_pct = round(float(tp2_pct), 2)

        rsi_val = c.get("rsi")
        mfi_val = c.get("mfi")
        stoch_k = c.get("stoch_k")
        stoch_d = c.get("stoch_d")
        macd_hist = c.get("macd_hist")

        # Stoch RSI Status
        if stoch_k is not None and stoch_d is not None:
            cross_str = "Bullish" if float(stoch_k) >= float(stoch_d) else "Bearish"
            stoch_str = f"{float(stoch_k):.1f} / {float(stoch_d):.1f} ({cross_str})"
        elif stoch_k is not None:
            stoch_str = f"%K: {float(stoch_k):.1f}"
        else:
            stoch_str = "Netral"

        # MACD Status
        if macd_hist is not None:
            macd_hist_f = float(macd_hist)
            if macd_hist_f > 0:
                macd_str = f"Bullish (+{macd_hist_f:.2f})"
            elif macd_hist_f < 0:
                macd_str = f"Bearish ({macd_hist_f:.2f})"
            else:
                macd_str = "Netral (0.0)"
        else:
            macd_str = "Netral"

        der_val = fund.get("der")
        der_str = f"{float(der_val):.2f}%" if (der_val is not None and der_val != "") else "-"

        roe_val = fund.get("roe")
        roe_str = f"{float(roe_val):.2f}%" if (roe_val is not None and roe_val != "") else "-"

        vol_ratio = float(c.get("vol_ratio") or bandar.get("vol_ratio", 1.0))

        row_data = [
            i,
            c.get("code", ""),
            c.get("name", ""),
            c.get("sector", ""),
            c.get("score", 0),
            c.get("setup_name") or c.get("setup_label") or c.get("setup_type", ""),
            c.get("risk_label") or c.get("risk_level", ""),
            close_price,
            entry,
            sl,
            f"-{sl_pct}%",
            tp1,
            f"+{tp1_pct}%",
            tp2,
            f"+{tp2_pct}%",
            tp.get("risk_reward") or tp.get("rr_ratio", "1 : 2.5"),
            bandar.get("label") or bandar.get("status_label") or bandar.get("status", "Netral"),
            round(vol_ratio, 2),
            round(float(rsi_val), 1) if rsi_val is not None else "-",
            round(float(mfi_val), 1) if mfi_val is not None else "-",
            stoch_str,
            macd_str,
            der_str,
            roe_str
        ]

        is_even = (i % 2 == 0)
        for col_idx, val in enumerate(row_data, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.font = normal_font
            cell.border = thin_border
            if is_even:
                cell.fill = zebra_fill

            # Alignments & Numbers
            if col_idx in [1, 2, 5, 7, 16, 17, 18, 19, 20, 21, 22, 23, 24]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif col_idx in [8, 9, 10, 12, 14]:
                cell.alignment = Alignment(horizontal="right", vertical="center")
                cell.number_format = '#,##0'
            elif col_idx in [11, 13, 15]:
                cell.alignment = Alignment(horizontal="right", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

        ws.row_dimensions[row_idx].height = 20

    # Auto-fit column widths
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.row < header_row:
                continue
            if cell.value:
                val_str = str(cell.value)
                if len(val_str) > max_len:
                    max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 11)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output


# Initialize database schema on module load
init_history_db()
